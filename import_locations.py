#!/usr/bin/env python3
"""Update directory.json's `sites` list from the "All Locations" tab of
Food Assistance Locations.xlsx.

Workflow:
  1. Edit "All Locations" in Food Assistance Locations.xlsx and save.
  2. python3 import_locations.py
  3. python3 generate.py   (rebuilds index.html + benefits.html)

Only `sites` is touched; `meta` and `programs` in directory.json are left as-is.
Existing rows are matched back to directory.json by the "ID (do not edit)"
column, so re-running this repeatedly produces a minimal diff. A row with a
blank ID is treated as a new location and gets an ID generated from its name.
A site whose ID is in directory.json but no longer present in the sheet is
treated as deleted and removed.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

BUILD = Path(__file__).parent
XLSX_PATH = BUILD / "Food Assistance Locations.xlsx"
DIRECTORY_PATH = BUILD / "directory.json"
SHEET_NAME = "All Locations"

CAT_LABELS = {
    "pantry": "Pantry",
    "soup_kitchen": "Soup Kitchen",
    "mini_pantry": "Mini-Pantry",
    "hot_meals": "Hot Meals",
    "market": "Market",
    "snap": "SNAP",
    "wic": "WIC",
    "senior": "Senior",
    "school": "School",
}
LABEL_TO_CAT = {v.lower(): k for k, v in CAT_LABELS.items()}

# header -> field name, in the same order build_xlsx.py writes them
COLUMNS = [
    ("ID (do not edit)", "id"),
    ("Name", "name"),
    ("Program Type(s)", "categories"),
    ("Town", "town"),
    ("Address", "address"),
    ("Latitude", "lat"),
    ("Longitude", "lng"),
    ("Phone", "phone"),
    ("Email", "email"),
    ("Website", "website"),
    ("Contact Person", "contactPerson"),
    ("Schedule", "schedule"),
    ("Days", "days"),
    ("Residents Served", "residentsServed"),
    ("ID Required", "idRequired"),
    ("Open to All", "openToAll"),
    ("Donations Accepted", "donationsAccepted"),
    ("Accepts (Markets)", "accepts"),
    ("Notes", "notes"),
]


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def to_bool(value, row_num, field):
    if isinstance(value, bool):
        return value
    if value is None:
        print(f"  warning: row {row_num} has a blank '{field}' — defaulting to False")
        return False
    s = str(value).strip().lower()
    if s in ("true", "yes", "y", "1"):
        return True
    if s in ("false", "no", "n", "0", ""):
        return False
    raise SystemExit(f"Row {row_num}: can't parse '{field}' value {value!r} as True/False")


def to_list(value):
    if value is None:
        return []
    return [part.strip() for part in str(value).split(",") if part.strip()]


def parse_categories(value, row_num):
    labels = to_list(value)
    cats = []
    for label in labels:
        cat = LABEL_TO_CAT.get(label.lower())
        if cat is None:
            valid = ", ".join(CAT_LABELS.values())
            raise SystemExit(
                f"Row {row_num}: unrecognized Program Type '{label}'. Valid values: {valid}"
            )
        cats.append(cat)
    return cats


def slugify(name):
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return slug or "location"


def unique_id(base, taken):
    candidate = base
    n = 2
    while candidate in taken:
        candidate = f"{base}-{n}"
        n += 1
    taken.add(candidate)
    return candidate


def main():
    if not XLSX_PATH.exists():
        raise SystemExit(f"Can't find {XLSX_PATH.name} next to import_locations.py")

    directory = json.loads(DIRECTORY_PATH.read_text(encoding="utf-8"))
    original_sites = directory["sites"]
    original_by_id = {s["id"]: s for s in original_sites}

    wb = load_workbook(XLSX_PATH, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"'{SHEET_NAME}' tab not found in {XLSX_PATH.name}")
    ws = wb[SHEET_NAME]

    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    field_by_col = {}
    for idx, header_text in enumerate(header):
        for label, field in COLUMNS:
            if header_text == label:
                field_by_col[idx] = field
                break
    missing = [label for label, _ in COLUMNS if label not in header]
    if missing:
        raise SystemExit(f"'{SHEET_NAME}' is missing expected column(s): {missing}")

    seen_ids = set(original_by_id.keys())
    rows_by_id = {}
    new_rows = []
    row_num = 1
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_num += 1
        raw = {}
        for idx, value in enumerate(row):
            field = field_by_col.get(idx)
            if field:
                raw[field] = value
        name = clean(raw.get("name"))
        if not name:
            continue  # skip blank rows

        site = {
            "id": clean(raw.get("id")),
            "name": name,
            "categories": parse_categories(raw.get("categories"), row_num),
            "town": clean(raw.get("town")),
            "address": clean(raw.get("address")),
            "phone": clean(raw.get("phone")),
            "email": clean(raw.get("email")),
            "website": clean(raw.get("website")),
            "contactPerson": clean(raw.get("contactPerson")),
            "residentsServed": clean(raw.get("residentsServed")),
            "idRequired": clean(raw.get("idRequired")) or "unknown",
            "schedule": clean(raw.get("schedule")),
            "days": to_list(raw.get("days")),
            "donationsAccepted": to_bool(raw.get("donationsAccepted"), row_num, "Donations Accepted"),
            "notes": clean(raw.get("notes")),
            "openToAll": to_bool(raw.get("openToAll"), row_num, "Open to All"),
            "lat": clean(raw.get("lat")),
            "lng": clean(raw.get("lng")),
        }
        accepts = to_list(raw.get("accepts"))
        if accepts:
            site["accepts"] = accepts

        site_id = site["id"]
        if site_id:
            if site_id in rows_by_id:
                raise SystemExit(f"Row {row_num}: duplicate ID '{site_id}' also used by another row")
            rows_by_id[site_id] = site
        else:
            new_rows.append(site)

    unknown_ids = set(rows_by_id) - set(original_by_id)
    if unknown_ids:
        raise SystemExit(
            "These IDs appear in the sheet but not in directory.json (typo, or "
            f"row copied from elsewhere?): {sorted(unknown_ids)}"
        )

    deleted_ids = set(original_by_id) - set(rows_by_id)

    updated_sites = []
    for site in original_sites:
        sid = site["id"]
        if sid in deleted_ids:
            print(f"  removed: {site['name']} ({sid})")
            continue
        edited = rows_by_id[sid]
        merged = dict(site)  # keep fields the sheet doesn't carry (kind, programId, contact, ...)
        merged.update(edited)
        merged["id"] = sid
        updated_sites.append(merged)

    for site in new_rows:
        new_id = unique_id(slugify(site["name"]), seen_ids)
        site["id"] = new_id
        site["kind"] = "site"
        print(f"  added: {site['name']} ({new_id})")
        updated_sites.append(site)

    directory["sites"] = updated_sites

    text = json.dumps(directory, indent=2, ensure_ascii=True)
    DIRECTORY_PATH.write_text(text, encoding="utf-8")
    print(
        f"Updated {DIRECTORY_PATH.name}: {len(updated_sites)} sites "
        f"({len(new_rows)} added, {len(deleted_ids)} removed, "
        f"{len(updated_sites) - len(new_rows)} kept/edited)."
    )
    print("Now run: python3 generate.py")


if __name__ == "__main__":
    sys.exit(main())
